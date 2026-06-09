#!/usr/bin/env python
"""Build a Schema After workbook draft from a GenBI Schema Before workbook."""

from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

for stream_name in ("stdout", "stderr"):
    stream = getattr(sys, stream_name, None)
    if hasattr(stream, "reconfigure"):
        stream.reconfigure(encoding="utf-8", errors="replace")

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover - environment guidance
    raise SystemExit(
        "缺少 openpyxl。請使用 Codex bundled Python，或先安裝 openpyxl 後再執行。"
    ) from exc


OUTPUT_HEADERS = [
    "schema",
    "table_name",
    "field",
    "table_description",
    "table_schema",
    "table_summary",
]
SCHEMA_SOURCE_DIR = Path("01-source") / "schema"
EXCEL_PATTERNS = ["*.xlsx", "*.xlsm", "*.xls"]


def latest_file(directory: Path, patterns: list[str]) -> Path | None:
    candidates: list[Path] = []
    if not directory.exists():
        return None
    for pattern in patterns:
        candidates.extend(
            path
            for path in directory.glob(pattern)
            if path.is_file() and not path.name.startswith("~$") and path.name != ".gitkeep"
        )
    if not candidates:
        return None
    return max(candidates, key=lambda path: path.stat().st_mtime)


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in {"nan", "none", "null"}:
        return ""
    return text


def normalize_header(value: Any) -> str:
    return cell_text(value).replace("\n", "").replace(" ", "").upper()


def find_header_map(ws: Any) -> dict[str, int]:
    headers: dict[str, int] = {}
    for cell in ws[1]:
        key = normalize_header(cell.value)
        if key:
            headers[key] = cell.column
    return headers


def first_existing(headers: dict[str, int], candidates: list[str]) -> int:
    for candidate in candidates:
        key = normalize_header(candidate)
        if key in headers:
            return headers[key]
        for header, column in headers.items():
            if key and (key in header or header in key):
                return column
    raise ValueError("找不到必要欄位：" + " / ".join(candidates))


def split_table_name(raw_name: str, default_schema: str = "public") -> tuple[str, str]:
    name = raw_name.strip().strip('"')
    if "." in name:
        schema, table = name.split(".", 1)
        return schema.strip('"').lower(), table.strip('"')
    return default_schema, name


def first_sentence(summary: str) -> str:
    if not summary:
        return ""
    candidates = [summary]
    for sep in ["，", ",", "。", "\n"]:
        idx = summary.find(sep)
        if idx >= 0:
            candidates.append(summary[:idx])
    return min(candidates, key=len).strip()


def quote_ident(identifier: str) -> str:
    return '"' + identifier.replace('"', '""') + '"'


def escape_comment(text: str) -> str:
    return text.replace("'", "''")


def build_comment(desc: str, notes: str) -> str:
    if desc and notes:
        return f"{desc}, 備註：{notes}"
    return desc or notes


def make_create_table(schema: str, table: str, columns: list[dict[str, str]]) -> str:
    lines = [f"CREATE TABLE {quote_ident(schema)}.{quote_ident(table)} ("]
    column_lines: list[str] = []
    for column in columns:
        name = column["name"]
        col_type = column["type"] or "TEXT"
        comment = build_comment(column["desc"], column["notes"])
        line = f"  {quote_ident(name)} {col_type}"
        if comment:
            line += f" COMMENT '{escape_comment(comment)}'"
        column_lines.append(line)
    lines.append(",\n".join(column_lines))
    lines.append(");")
    return "\n".join(lines)


def load_table_rows(wb: Any) -> list[dict[str, str]]:
    if "資料表總表" not in wb.sheetnames:
        raise ValueError("輸入檔缺少 sheet：資料表總表")
    ws = wb["資料表總表"]
    headers = find_header_map(ws)
    scene_col = first_existing(headers, ["場景"])
    table_col = first_existing(headers, ["資料表名稱"])
    summary_col = first_existing(headers, ["資料表說明"])

    rows: list[dict[str, str]] = []
    for row_idx in range(2, ws.max_row + 1):
        raw_table = cell_text(ws.cell(row_idx, table_col).value)
        if not raw_table:
            continue
        schema, table = split_table_name(raw_table)
        summary = cell_text(ws.cell(row_idx, summary_col).value)
        rows.append(
            {
                "schema": schema,
                "table_name": table,
                "field": cell_text(ws.cell(row_idx, scene_col).value),
                "table_description": first_sentence(summary),
                "table_summary": summary,
            }
        )
    return rows


def load_columns_by_table(wb: Any) -> dict[str, list[dict[str, str]]]:
    if "表欄位" not in wb.sheetnames:
        raise ValueError("輸入檔缺少 sheet：表欄位")
    ws = wb["表欄位"]
    headers = find_header_map(ws)
    table_col = first_existing(headers, ["資料表名稱"])
    num_col = first_existing(headers, ["欄位標號", "DW_COLUMN_NUM"])
    name_col = first_existing(headers, ["欄位名稱", "DW_COLUMN_NAME"])
    type_col = first_existing(headers, ["欄位格式", "DW_COLUMN_FORMAT"])
    desc_col = first_existing(headers, ["欄位描述", "COLUMN_DESC"])
    notes_col = first_existing(headers, ["欄位備註", "COLUMN_NOTES"])

    grouped: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row_idx in range(2, ws.max_row + 1):
        raw_table = cell_text(ws.cell(row_idx, table_col).value)
        name = cell_text(ws.cell(row_idx, name_col).value)
        if not raw_table or not name:
            continue
        _, table = split_table_name(raw_table)
        grouped[table].append(
            {
                "row_idx": str(row_idx),
                "num": cell_text(ws.cell(row_idx, num_col).value),
                "name": name,
                "type": cell_text(ws.cell(row_idx, type_col).value),
                "desc": cell_text(ws.cell(row_idx, desc_col).value),
                "notes": cell_text(ws.cell(row_idx, notes_col).value),
            }
        )

    def sort_key(column: dict[str, str]) -> tuple[int, str]:
        try:
            return int(float(column["num"])), column["row_idx"]
        except ValueError:
            return 999999, column["row_idx"]

    for columns in grouped.values():
        columns.sort(key=sort_key)
    return grouped


def autosize(ws: Any) -> None:
    widths = {
        "A": 16,
        "B": 36,
        "C": 14,
        "D": 34,
        "E": 90,
        "F": 90,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")


def build_schema_after(input_path: Path, output_path: Path) -> None:
    if not input_path.exists():
        raise FileNotFoundError(
            f"找不到輸入檔：{input_path}。請將正式 schema workbook 放到 01-source/schema/，"
            "或使用 --input 指定其他檔案。00-examples/ 僅作格式參考。"
        )
    wb_in = openpyxl.load_workbook(input_path, read_only=True, data_only=True)
    table_rows = load_table_rows(wb_in)
    columns_by_table = load_columns_by_table(wb_in)

    wb_out = openpyxl.Workbook()
    ws = wb_out.active
    ws.title = "Sheet1"
    ws.append(OUTPUT_HEADERS)

    for table_row in table_rows:
        schema = table_row["schema"]
        table = table_row["table_name"]
        table_schema = make_create_table(schema, table, columns_by_table.get(table, []))
        ws.append(
            [
                schema,
                table,
                table_row["field"],
                table_row["table_description"],
                table_schema,
                table_row["table_summary"],
            ]
        )

    autosize(ws)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb_out.save(output_path)
    except PermissionError as exc:
        raise PermissionError(
            f"無法寫入輸出檔：{output_path}。請確認檔案沒有被 Excel 或其他程式開啟，"
            "或使用 --output 指定新的輸出檔名。"
        ) from exc


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input",
        type=Path,
        default=None,
        help="Schema Before workbook. Defaults to latest Excel under 01-source/schema/.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("03-final-exports") / "schema" / "Schema__after.xlsx",
        help="Schema After output workbook",
    )
    args = parser.parse_args()
    input_path = args.input or latest_file(SCHEMA_SOURCE_DIR, EXCEL_PATTERNS)
    if input_path is None:
        print(
            "找不到 schema Excel。請將檔案放到 01-source/schema/，"
            "或使用 --input 指定檔案。檔名可自訂。"
        )
        return 1

    try:
        build_schema_after(input_path, args.output)
    except FileNotFoundError as exc:
        print(exc)
        return 1
    except PermissionError as exc:
        print(exc)
        return 1
    print(f"使用輸入檔：{input_path}")
    print(f"已產生 Schema After 草稿：{args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

