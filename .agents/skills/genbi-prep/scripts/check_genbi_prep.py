#!/usr/bin/env python
"""Inspect GenBI prep workbooks and print actionable TODOs."""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

for stream_name in ("stdout", "stderr"):
    stream = getattr(sys, stream_name, None)
    if hasattr(stream, "reconfigure"):
        stream.reconfigure(encoding="utf-8", errors="replace")

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover - environment guidance
    raise SystemExit(
        "缺少 openpyxl。請使用 Codex bundled Python，或先安裝 openpyxl 後再執行。"
    ) from exc


SCHEMA_DIR_REL = Path("01-source") / "schema"
KNOWLEDGE_DIR_REL = Path("01-source") / "knowledge"
INTERVIEW_REL = Path("01-source") / "interview"
GOLDEN_DIR_REL = Path("01-source") / "golden"
WORKBENCH_ROOT_REL = Path("02-workbench")
EXAMPLE_SCHEMA_REL = Path("00-examples") / "01-source" / "schema" / "Schema__before__example.xlsx"
EXAMPLE_KNOWLEDGE_REL = Path("00-examples") / "01-source" / "knowledge" / "知識典__example.xlsx"
EXAMPLE_GOLDEN_REL = Path("00-examples") / "01-source" / "golden" / "golden__dataset__example.xlsx"

MIN_DESC_CHARS = 8
MIN_NOTE_CHARS = 16
MAX_EXAMPLES = 20
EXCEL_PATTERNS = ["*.xlsx", "*.xlsm", "*.xls"]


@dataclass
class Finding:
    severity: str
    area: str
    message: str


def latest_file(directory: Path, patterns: list[str]) -> Path | None:
    candidates: list[Path] = []
    if not directory.exists():
        return None
    for pattern in patterns:
        candidates.extend(
            path
            for path in directory.glob(pattern)
            if path.is_file() and not path.name.startswith("~$") and path.name != ".gitkeep"
        )
    if not candidates:
        return None
    return max(candidates, key=lambda path: path.stat().st_mtime)


def candidate_summary(directory: Path, patterns: list[str], limit: int = 5) -> str:
    candidates: list[Path] = []
    if directory.exists():
        for pattern in patterns:
            candidates.extend(
                path
                for path in directory.glob(pattern)
                if path.is_file() and not path.name.startswith("~$") and path.name != ".gitkeep"
            )
    if not candidates:
        return "無候選檔"
    ordered = sorted(candidates, key=lambda path: path.stat().st_mtime, reverse=True)[:limit]
    return "; ".join(path.name for path in ordered)


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in {"nan", "none", "null"}:
        return ""
    return text


def normalize_header(value: Any) -> str:
    return cell_text(value).replace("\n", "").replace(" ", "").upper()


def find_header_map(ws: Any) -> dict[str, int]:
    headers: dict[str, int] = {}
    for cell in ws[1]:
        key = normalize_header(cell.value)
        if key:
            headers[key] = cell.column
    return headers


def first_existing(headers: dict[str, int], candidates: list[str]) -> int | None:
    for candidate in candidates:
        key = normalize_header(candidate)
        if key in headers:
            return headers[key]
        for header, column in headers.items():
            if key and (key in header or header in key):
                return column
    return None


def is_code_like(column_name: str, column_format: str, desc: str) -> bool:
    haystack = f"{column_name} {column_format} {desc}".lower()
    hints = [
        "code",
        "seq",
        "status",
        "type",
        "flag",
        "level",
        "category",
        "代碼",
        "狀態",
        "類別",
        "分類",
        "等級",
        "是否",
        "註記",
    ]
    return any(hint in haystack for hint in hints)


def inspect_schema(path: Path | None, source_dir: Path | None = None) -> list[Finding]:
    findings: list[Finding] = []
    if path is None:
        return [
            Finding(
                "ERROR",
                "Schema",
                f"找不到 schema Excel；請將檔案放到 {source_dir}，或用 --schema 指定檔案。",
            )
        ]
    if not path.exists():
        return [Finding("ERROR", "Schema", f"找不到檔案：{path}")]

    if source_dir:
        findings.append(
            Finding(
                "INFO",
                "Schema",
                f"本次檢查使用最新 schema 候選檔：{path}。候選檔：{candidate_summary(source_dir, EXCEL_PATTERNS)}",
            )
        )

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    required_sheets = ["資料表總表", "表欄位"]
    for sheet in required_sheets:
        if sheet not in wb.sheetnames:
            findings.append(Finding("ERROR", "Schema", f"{path.name} 缺少 sheet：{sheet}"))
    if any(f.severity == "ERROR" for f in findings):
        return findings

    table_ws = wb["資料表總表"]
    table_headers = find_header_map(table_ws)
    scene_col = first_existing(table_headers, ["場景"])
    table_name_col = first_existing(table_headers, ["資料表名稱"])
    table_desc_col = first_existing(table_headers, ["資料表說明"])

    missing_table_desc: list[str] = []
    total_table_names: set[str] = set()
    for row_idx in range(2, table_ws.max_row + 1):
        table_name = cell_text(table_ws.cell(row_idx, table_name_col).value) if table_name_col else ""
        if not table_name:
            continue
        total_table_names.add(table_name.split(".")[-1])
        table_desc = cell_text(table_ws.cell(row_idx, table_desc_col).value) if table_desc_col else ""
        scene = cell_text(table_ws.cell(row_idx, scene_col).value) if scene_col else ""
        if len(table_desc) < MIN_DESC_CHARS:
            missing_table_desc.append(f"row {row_idx} table={table_name} field={scene or '-'}")

    if missing_table_desc:
        findings.append(
            Finding(
                "TODO",
                "資料表總表",
                "資料表說明空白或過短：" + "; ".join(missing_table_desc[:MAX_EXAMPLES]),
            )
        )

    field_ws = wb["表欄位"]
    field_headers = find_header_map(field_ws)
    table_col = first_existing(field_headers, ["資料表名稱"])
    col_num_col = first_existing(field_headers, ["欄位標號", "DW_COLUMN_NUM"])
    name_col = first_existing(field_headers, ["欄位名稱", "DW_COLUMN_NAME"])
    format_col = first_existing(field_headers, ["欄位格式", "DW_COLUMN_FORMAT"])
    desc_col = first_existing(field_headers, ["欄位描述", "COLUMN_DESC"])
    notes_col = first_existing(field_headers, ["欄位備註", "COLUMN_NOTES"])
    common_col = first_existing(field_headers, ["是否為常用欄位"])
    key_col = first_existing(field_headers, ["是否為KEY值", "是否為key值"])

    required_cols = {
        "資料表名稱": table_col,
        "欄位名稱": name_col,
        "欄位格式": format_col,
        "欄位描述": desc_col,
        "欄位備註": notes_col,
    }
    missing_cols = [name for name, col in required_cols.items() if col is None]
    if missing_cols:
        return [
            Finding(
                "ERROR",
                "表欄位",
                f"{path.name} 的 `表欄位` 缺少必要欄位：" + ", ".join(missing_cols),
            )
        ]

    missing_desc: list[str] = []
    thin_desc: list[str] = []
    missing_notes_priority: list[str] = []
    field_table_names: set[str] = set()
    total_fields = 0
    desc_ready = 0
    notes_ready = 0

    for row_idx in range(2, field_ws.max_row + 1):
        table_name = cell_text(field_ws.cell(row_idx, table_col).value)
        col_name = cell_text(field_ws.cell(row_idx, name_col).value)
        if not table_name and not col_name:
            continue
        if table_name:
            field_table_names.add(table_name.split(".")[-1])
        total_fields += 1
        col_num = cell_text(field_ws.cell(row_idx, col_num_col).value) if col_num_col else ""
        col_format = cell_text(field_ws.cell(row_idx, format_col).value)
        desc = cell_text(field_ws.cell(row_idx, desc_col).value)
        notes = cell_text(field_ws.cell(row_idx, notes_col).value)
        common = cell_text(field_ws.cell(row_idx, common_col).value) if common_col else ""
        key = cell_text(field_ws.cell(row_idx, key_col).value) if key_col else ""
        label = f"row {row_idx} {table_name}.{col_name}"
        if col_num:
            label += f" #{col_num}"

        if not desc:
            missing_desc.append(label)
        elif len(desc) < MIN_DESC_CHARS:
            thin_desc.append(label)
        else:
            desc_ready += 1

        priority_note = common == "1" or key == "1" or is_code_like(col_name, col_format, desc)
        if notes:
            notes_ready += 1
        elif priority_note:
            missing_notes_priority.append(label)

    findings.append(
        Finding(
            "INFO",
            "表欄位",
            f"欄位總數 {total_fields}；描述看起來可用 {desc_ready}；備註已填 {notes_ready}。",
        )
    )

    if missing_desc:
        findings.append(
            Finding(
                "TODO",
                "表欄位",
                "欄位描述空白：" + "; ".join(missing_desc[:MAX_EXAMPLES]),
            )
        )
    if thin_desc:
        findings.append(
            Finding(
                "TODO",
                "表欄位",
                "欄位描述可能過短，需要 enrich：" + "; ".join(thin_desc[:MAX_EXAMPLES]),
            )
        )
    if missing_notes_priority:
        findings.append(
            Finding(
                "TODO",
                "表欄位",
                "常用/key/代碼類欄位缺少備註，需追問 SQL 使用規則或代碼意涵："
                + "; ".join(missing_notes_priority[:MAX_EXAMPLES]),
            )
        )

    missing_from_fields = sorted(total_table_names - field_table_names)
    if missing_from_fields:
        findings.append(
            Finding("TODO", "資料表總表", "資料表總表有列出但表欄位找不到欄位：" + ", ".join(missing_from_fields))
        )

    return findings


def inspect_examples(root: Path) -> list[Finding]:
    expected = [
        EXAMPLE_SCHEMA_REL,
        Path("00-examples") / "02-workbench" / "schema" / "20260525-1430_schema_enriched_v01_example_draft.xlsx",
        Path("00-examples") / "02-workbench" / "schema" / "CHANGELOG.md",
        Path("00-examples") / "03-final-exports" / "schema" / "Schema__after__example.xlsx",
        EXAMPLE_KNOWLEDGE_REL,
        EXAMPLE_GOLDEN_REL,
    ]
    missing = [str(path) for path in expected if not (root / path).exists()]
    if missing:
        return [Finding("TODO", "Examples", "範例檔缺少：" + "; ".join(missing))]
    return [Finding("INFO", "Examples", "範例檔已集中在 00-examples/，僅作格式參考，不作為工作檔。")]


def inspect_knowledge(path: Path | None, source_dir: Path | None = None) -> list[Finding]:
    if path is None:
        return [
            Finding(
                "TODO",
                "知識典",
                f"找不到知識典 Excel；請將檔案放到 {source_dir}，或用 --knowledge 指定檔案。",
            )
        ]
    if not path.exists():
        return [Finding("TODO", "知識典", f"找不到檔案：{path}")]
    prefix: list[Finding] = []
    if source_dir:
        prefix.append(
            Finding(
                "INFO",
                "知識典",
                f"本次檢查使用最新知識典候選檔：{path}。候選檔：{candidate_summary(source_dir, EXCEL_PATTERNS)}",
            )
        )
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "知識典" not in wb.sheetnames:
        return [Finding("ERROR", "知識典", f"{path.name} 缺少 sheet：知識典")]
    ws = wb["知識典"]
    headers = find_header_map(ws)
    field_col = first_existing(headers, ["field"])
    keyword_col = first_existing(headers, ["關鍵字"])
    content_col = first_existing(headers, ["內容"])
    if not all([field_col, keyword_col, content_col]):
        return [Finding("ERROR", "知識典", "`知識典` sheet 應包含 field、關鍵字、內容。")]

    filled = 0
    missing_content: list[str] = []
    for row_idx in range(2, ws.max_row + 1):
        field = cell_text(ws.cell(row_idx, field_col).value)
        keyword = cell_text(ws.cell(row_idx, keyword_col).value)
        content = cell_text(ws.cell(row_idx, content_col).value)
        if not field and not keyword and not content:
            continue
        if keyword and content:
            filled += 1
        elif keyword and not content:
            missing_content.append(f"row {row_idx} keyword={keyword}")

    findings = prefix + [Finding("INFO", "知識典", f"目前已有 {filled} 筆看起來可用的知識典項目。")]
    if filled == 0:
        findings.append(Finding("TODO", "知識典", "知識典目前沒有可用項目；訪談時若發現共通口徑，請新增。"))
    if missing_content:
        findings.append(Finding("TODO", "知識典", "關鍵字有填但內容空白：" + "; ".join(missing_content[:MAX_EXAMPLES])))
    return findings


def inspect_interview(path: Path) -> list[Finding]:
    if not path.exists():
        return [Finding("TODO", "使用者需求訪談", f"找不到資料夾：{path}")]
    files = [
        p
        for p in path.rglob("*")
        if p.is_file() and not p.name.startswith("~$") and p.name != ".gitkeep"
    ]
    if not files:
        return [
            Finding(
                "TODO",
                "使用者需求訪談",
                "訪談資料夾目前沒有檔案；請記錄 agent 目標、常見問題、常用查詢與輸出期待。",
            )
        ]
    return [Finding("INFO", "使用者需求訪談", f"訪談資料夾已有 {len(files)} 個檔案。")]


def inspect_workbench_changelogs(root: Path) -> list[Finding]:
    findings: list[Finding] = []
    for artifact in ["schema", "knowledge", "golden", "interview", "reports"]:
        directory = root / WORKBENCH_ROOT_REL / artifact
        changelog = directory / "CHANGELOG.md"
        has_files = False
        if directory.exists():
            has_files = any(
                path.is_file()
                and path.name not in {".gitkeep", "CHANGELOG.md"}
                and not path.name.startswith("~$")
                for path in directory.iterdir()
            )
        if changelog.exists():
            findings.append(Finding("INFO", "Workbench 版本控制", f"{artifact} 已建立 changelog：{changelog}"))
        elif has_files:
            findings.append(
                Finding(
                    "TODO",
                    "Workbench 版本控制",
                    f"{artifact} 已有工作檔但找不到 {changelog}；請在該 artifact 資料夾內建立 changelog。",
                )
            )
    if not findings:
        findings.append(
            Finding(
                "INFO",
                "Workbench 版本控制",
                "尚未建立 artifact changelog；開始產生 workbench 檔後，請在同資料夾建立 CHANGELOG.md。",
            )
        )
    return findings


def inspect_golden(path: Path | None, source_dir: Path | None = None) -> list[Finding]:
    if path is None:
        return [
            Finding(
                "INFO",
                "Golden Dataset",
                f"尚未提供 golden dataset；這通常會在訪談與 schema enrich 後建立。若已有檔案可放到 {source_dir}，或用 --golden 指定。",
            )
        ]
    if not path.exists():
        return [Finding("TODO", "Golden Dataset", f"找不到檔案：{path}")]
    prefix: list[Finding] = []
    if source_dir:
        prefix.append(
            Finding(
                "INFO",
                "Golden Dataset",
                f"本次檢查使用最新 golden 候選檔：{path}。候選檔：{candidate_summary(source_dir, EXCEL_PATTERNS)}",
            )
        )
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "evaluation_set" not in wb.sheetnames:
        return [Finding("ERROR", "Golden Dataset", f"{path.name} 缺少 sheet：evaluation_set")]
    ws = wb["evaluation_set"]
    headers = find_header_map(ws)
    question_col = first_existing(headers, ["question"])
    sql_col = first_existing(headers, ["correct_postgres_sql"])
    if not question_col or not sql_col:
        return [Finding("ERROR", "Golden Dataset", "`evaluation_set` 應包含 question 與 correct_postgres_sql。")]

    complete = 0
    missing_question: list[str] = []
    missing_sql: list[str] = []
    for row_idx in range(2, ws.max_row + 1):
        question = cell_text(ws.cell(row_idx, question_col).value)
        sql = cell_text(ws.cell(row_idx, sql_col).value)
        if not question and not sql:
            continue
        if question and sql:
            complete += 1
        elif not question:
            missing_question.append(f"row {row_idx}")
        elif not sql:
            missing_sql.append(f"row {row_idx}: {question[:40]}")

    findings = prefix + [Finding("INFO", "Golden Dataset", f"目前已有 {complete} 筆 question + SQL 成對資料。")]
    if complete == 0:
        findings.append(Finding("TODO", "Golden Dataset", "尚未建立可用 Q&A；請先設計高頻問題與正確 PostgreSQL。"))
    if missing_question:
        findings.append(Finding("TODO", "Golden Dataset", "SQL 有填但 question 空白：" + "; ".join(missing_question[:MAX_EXAMPLES])))
    if missing_sql:
        findings.append(Finding("TODO", "Golden Dataset", "question 有填但 correct_postgres_sql 空白：" + "; ".join(missing_sql[:MAX_EXAMPLES])))
    return findings


def print_findings(findings: list[Finding]) -> None:
    grouped: dict[str, list[Finding]] = {}
    for finding in findings:
        grouped.setdefault(finding.area, []).append(finding)

    print("# GenBI 前置文件檢查結果\n")
    for area, items in grouped.items():
        print(f"## {area}")
        for item in items:
            print(f"- [{item.severity}] {item.message}")
        print()

    todos = [f for f in findings if f.severity in {"ERROR", "TODO"}]
    errors = [f for f in findings if f.severity == "ERROR"]
    if todos:
        print("## 建議下一步")
        if errors:
            print("- 先針對 [ERROR] 修正檔案或 sheet 結構。")
        missing_source_files = [
            f
            for f in todos
            if "source" in f.message and ("找不到" in f.message)
        ]
        if missing_source_files:
            print("- 將使用者正式提供的 Excel 放到 01-source/；00-examples/ 只當格式參考。")
        print("- 修改任何 workbench 檔案前，先執行 list_active_candidates.py 並與使用者確認 active file。")
        print("- 確認 active file 後，先用 version_workbench_artifact.py 建立新版本；只修改新版本檔。")
        print("- 再挑 1-2 張最重要資料表，向使用者追問欄位描述、代碼意涵、常用 SQL 規則。")
        print("- 若答案是跨欄位共通口徑，請寫入知識典；若是單欄位規則，請寫入欄位備註。")
    else:
        print("## 建議下一步")
        print("- 文件基本完整；可以開始產生 Schema After 或擴充 golden Q&A。")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path("."), help="Workspace root")
    parser.add_argument("--schema", type=Path, default=None, help="Schema before workbook")
    parser.add_argument("--knowledge", type=Path, default=None, help="Knowledge workbook")
    parser.add_argument("--golden", type=Path, default=None, help="Golden dataset workbook")
    args = parser.parse_args()

    root = args.root.resolve()
    schema_dir = root / SCHEMA_DIR_REL
    knowledge_dir = root / KNOWLEDGE_DIR_REL
    golden_dir = root / GOLDEN_DIR_REL
    schema = args.schema or latest_file(schema_dir, EXCEL_PATTERNS)
    knowledge = args.knowledge or latest_file(knowledge_dir, EXCEL_PATTERNS)
    golden = args.golden or latest_file(golden_dir, EXCEL_PATTERNS)
    interview = root / INTERVIEW_REL

    findings: list[Finding] = []
    findings.extend(inspect_examples(root))
    findings.extend(inspect_schema(schema, None if args.schema else schema_dir))
    findings.extend(inspect_knowledge(knowledge, None if args.knowledge else knowledge_dir))
    findings.extend(inspect_interview(interview))
    findings.extend(inspect_workbench_changelogs(root))
    findings.extend(inspect_golden(golden, None if args.golden else golden_dir))
    print_findings(findings)
    return 1 if any(f.severity == "ERROR" for f in findings) else 0


if __name__ == "__main__":
    raise SystemExit(main())


